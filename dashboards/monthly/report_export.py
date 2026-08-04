"""월간 대시보드 내용을 세로형 Excel·PDF 보고서로 변환한다."""

from io import BytesIO

import pandas as pd


PURPLE = "6C5CE7"
NAVY = "252A4A"
LIGHT_PURPLE = "EEEAFE"
LIGHT_GRAY = "F3F5F9"
WHITE = "FFFFFF"


def _display_value(value):
    if pd.isna(value):
        return "-"
    if isinstance(value, float):
        return f"{value:,.1f}"
    return value


def create_excel_report(report):
    """요약 보고서와 상세 시트를 포함한 세로형 xlsx 바이트를 반환한다."""
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "월간보고서"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.3
    sheet.page_margins.right = 0.3
    sheet.column_dimensions["A"].width = 23
    sheet.column_dimensions["B"].width = 18
    sheet.column_dimensions["C"].width = 18
    sheet.column_dimensions["D"].width = 18

    thin = Side(style="thin", color="D9DDEA")
    row = 1
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    cell = sheet.cell(row, 1, f"잇뉴 {report['month_label']} 월간 분석 보고서")
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(color=WHITE, bold=True, size=20)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[row].height = 42
    row += 2
    sheet.cell(row, 1, "조회 기간").font = Font(bold=True, color=PURPLE)
    sheet.cell(row, 2, report["period_label"])
    sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    row += 2

    def section(title, dataframe, chart=False):
        nonlocal row
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        heading = sheet.cell(row, 1, title)
        heading.fill = PatternFill("solid", fgColor=LIGHT_PURPLE)
        heading.font = Font(bold=True, color=NAVY, size=13)
        heading.alignment = Alignment(vertical="center")
        sheet.row_dimensions[row].height = 25
        row += 1
        start_header = row
        for column_index, column in enumerate(dataframe.columns, 1):
            header = sheet.cell(row, column_index, str(column))
            header.fill = PatternFill("solid", fgColor=NAVY)
            header.font = Font(bold=True, color=WHITE)
            header.alignment = Alignment(horizontal="center")
            header.border = Border(bottom=thin)
        for values in dataframe.itertuples(index=False, name=None):
            row += 1
            for column_index, value in enumerate(values, 1):
                body = sheet.cell(row, column_index, _display_value(value))
                body.fill = PatternFill("solid", fgColor=WHITE if row % 2 else LIGHT_GRAY)
                body.border = Border(bottom=thin)
                body.alignment = Alignment(horizontal="left" if column_index == 1 else "right")
        if chart and len(dataframe) and len(dataframe.columns) >= 2:
            chart_object = BarChart()
            chart_object.type = "bar"
            chart_object.style = 10
            chart_object.title = title
            chart_object.height = 6.2
            chart_object.width = 13
            chart_object.add_data(
                Reference(sheet, min_col=2, min_row=start_header, max_row=row), titles_from_data=True
            )
            chart_object.set_categories(Reference(sheet, min_col=1, min_row=start_header + 1, max_row=row))
            chart_object.legend = None
            sheet.add_chart(chart_object, f"A{row + 2}")
            row += 13
        row += 2

    section("1. 월간 경영 지표", report["kpi_table"])
    section("2. 당월 주요 이슈·조치 내역·결과", report["issue_table"])
    section("3. 익월 계획", report["next_plan_table"])
    section("4. 협력사 요청 사항", report["partner_request_table"])
    section("5. 회원 현황", report["member_table"], chart=True)
    section("6. 품목별 배송 현황", report["product_table"].head(20), chart=True)

    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    heading = sheet.cell(row, 1, "7. 월간 자동 요약")
    heading.fill = PatternFill("solid", fgColor=LIGHT_PURPLE)
    heading.font = Font(bold=True, color=NAVY, size=13)
    row += 1
    for summary in report["summary_lines"]:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        cell = sheet.cell(row, 1, f"• {summary}")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[row].height = 28
        row += 1
    sheet.print_area = f"A1:D{row}"

    for title, dataframe in [
        ("품목상세", report["product_table"]),
        ("회원현황", report["member_table"]),
        ("주요이슈", report["issue_table"]),
    ]:
        detail_sheet = workbook.create_sheet(title)
        detail_sheet.sheet_view.showGridLines = False
        detail_sheet.page_setup.orientation = "portrait"
        detail_sheet.page_setup.fitToWidth = 1
        for column_index, column in enumerate(dataframe.columns, 1):
            cell = detail_sheet.cell(1, column_index, str(column))
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(bold=True, color=WHITE)
            detail_sheet.column_dimensions[cell.column_letter].width = 24
        for row_index, values in enumerate(dataframe.itertuples(index=False, name=None), 2):
            for column_index, value in enumerate(values, 1):
                detail_sheet.cell(row_index, column_index, _display_value(value))
        detail_sheet.freeze_panes = "A2"
        detail_sheet.auto_filter.ref = detail_sheet.dimensions

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def create_pdf_report(report):
    """A4 세로형 PDF 보고서 바이트를 반환한다."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    pdfmetrics.registerFont(UnicodeCIDFont("HYSMyeongJo-Medium"))
    output = BytesIO()
    document = SimpleDocTemplate(
        output, pagesize=A4, rightMargin=14 * mm, leftMargin=14 * mm,
        topMargin=14 * mm, bottomMargin=14 * mm,
        title=f"잇뉴 {report['month_label']} 월간 분석 보고서",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "KoreanTitle", parent=styles["Title"], fontName="HYSMyeongJo-Medium",
        fontSize=20, leading=28, textColor=colors.HexColor("#252A4A"), alignment=TA_CENTER,
    )
    heading_style = ParagraphStyle(
        "KoreanHeading", parent=styles["Heading2"], fontName="HYSMyeongJo-Medium",
        fontSize=13, leading=19, textColor=colors.HexColor("#4E3DBA"), spaceBefore=10, spaceAfter=7,
    )
    body_style = ParagraphStyle(
        "KoreanBody", parent=styles["BodyText"], fontName="HYSMyeongJo-Medium",
        fontSize=8.5, leading=13, textColor=colors.HexColor("#303445"),
    )

    story = [
        Paragraph(f"잇뉴 {report['month_label']} 월간 분석 보고서", title_style),
        Spacer(1, 5 * mm),
        Paragraph(f"조회 기간: {report['period_label']}", body_style),
        Paragraph("본 보고서는 대시보드의 비민감 집계 정보만 포함합니다.", body_style),
        Spacer(1, 5 * mm),
    ]

    def add_table(title, dataframe, limit=None):
        story.append(Paragraph(title, heading_style))
        shown = dataframe.head(limit) if limit else dataframe
        table_data = [[Paragraph(str(column), body_style) for column in shown.columns]]
        for values in shown.itertuples(index=False, name=None):
            table_data.append([Paragraph(str(_display_value(value)), body_style) for value in values])
        available_width = A4[0] - 28 * mm
        table = Table(
            table_data,
            colWidths=[available_width / max(len(shown.columns), 1)] * max(len(shown.columns), 1),
            repeatRows=1,
            hAlign="LEFT",
        )
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#252A4A")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, -1), "HYSMyeongJo-Medium"),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D9DDEA")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F5F9")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        story.extend([table, Spacer(1, 3 * mm)])

    add_table("1. 월간 경영 지표", report["kpi_table"])
    add_table("2. 당월 주요 이슈·조치 내역·결과", report["issue_table"])
    add_table("3. 익월 계획", report["next_plan_table"])
    add_table("4. 협력사 요청 사항", report["partner_request_table"])
    story.append(PageBreak())
    add_table("5. 회원 현황", report["member_table"])
    add_table("6. 품목별 배송 현황 (상위 20개)", report["product_table"], 20)
    story.append(Paragraph("7. 월간 자동 요약", heading_style))
    for summary in report["summary_lines"]:
        story.append(Paragraph(f"• {summary}", body_style))
        story.append(Spacer(1, 1.5 * mm))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("HYSMyeongJo-Medium", 7)
        canvas.setFillColor(colors.HexColor("#7A8093"))
        canvas.drawString(14 * mm, 8 * mm, "잇뉴 통합 분석 플랫폼")
        canvas.drawRightString(A4[0] - 14 * mm, 8 * mm, f"{doc.page} 페이지")
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()
